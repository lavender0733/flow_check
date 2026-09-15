# -*- coding: utf-8 -*-
"""
年度檢索統計獨立程式
====================

啟動：python program/yearly_stats/yearly_server.py
開啟：http://127.0.0.1:8766

統計來源：
  - data/620/：.xlsx 的「月份追蹤」工作表歷史表，從 2026-01 起加總「檢索次數」。
  - data/630/：每月 .xls 的「全部」工作表最下方「合計」列，讀取「檢索次數」。
  - data/680/：.xlsx 的「統計表」最下方「累計／總計」列，讀取「檢索次數」。

620／680 資料夾裡常同時留著舊檔，挑檔是比「資料實際涵蓋到的月份」而不是檔案時間 ——
Google Drive 同步會重寫 mtime，靠時間挑會在別台機器挑到舊檔。
"""

import json
import re
import sys
import threading
import webbrowser
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    import xlrd
except Exception:
    xlrd = None

sys.dont_write_bytecode = True

# 把 program/ 放上 sys.path，讓 core / gpss / common / yearly_stats 都能以套件匯入。
# 這是全專案唯一動 sys.path 的地方（每個進入點各一行），取代原本散在模組內的路徑魔法。
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

BASE = Path(__file__).resolve().parents[2]
DATA_DIR = BASE / "data"
TARGETS_PATH = DATA_DIR / "config" / "targets.json"
WEB_DIR = Path(__file__).resolve().with_name("web")
CACHE_PATH = DATA_DIR / "年度統計快取.json"
# 解析結果的欄位一改（例：680 開始附上各月份），舊快取的形狀就對不上；把版本併進簽章，
# 改版後第一次統計會自動重讀，不必手動刪快取檔。
CACHE_VERSION = 2
PORT = 8766
_CACHE_LOCK = threading.Lock()


def as_int(value) -> int:
    """把 Excel 數字或帶千分位的字串轉成整數。"""
    if value is None or isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return int(round(value))
    text = re.sub(r"[^0-9.-]", "", str(value))
    try:
        return int(round(float(text))) if text else 0
    except ValueError:
        return 0


def file_signature(path: Path) -> str:
    """檔案大小與修改時間組成的簽章；內容不變時不必再次開啟 Excel。"""
    try:
        stat = path.stat()
        return f"{stat.st_size}:{stat.st_mtime_ns}:v{CACHE_VERSION}"
    except OSError:
        return ""


def load_cache() -> dict:
    """讀取跨次啟動仍可使用的統計快取；壞檔安全退回空快取。"""
    try:
        data = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def save_cache(cache: dict) -> None:
    """以暫存檔原子替換快取，避免同步／中斷時留下半截 JSON。"""
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp = CACHE_PATH.with_suffix(".tmp")
    temp.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(CACHE_PATH)


def cached_parse(cache: dict, path: Path, parser):
    """簽章命中就回快取；檔案更新時才呼叫 parser 重新讀取 Excel。"""
    key = str(path.resolve())
    signature = file_signature(path)
    cached = cache.get(key)
    if cached and cached.get("signature") == signature:
        return cached.get("data"), True
    data = parser(path)
    cache[key] = {"signature": signature, "data": data}
    return data, False


def load_annual_target() -> int:
    try:
        value = int(json.loads(TARGETS_PATH.read_text(encoding="utf-8")).get("annual_target", 0) or 0)
        return max(0, value)
    except Exception:
        return 0


def month_key(path: Path):
    match = re.search(r"[（(]\s*(\d{1,2})\s*月\s*[)）]", path.name)
    return int(match.group(1)) if match else 99


def latest_data_month(data) -> int:
    """解析結果中「真的有數字」的最後一個月份（YYYYMM）；沒有就回 0。
    月報表常預先排好整年的月份列，只看月份欄會把還沒發生的月份也算進去。"""
    months = (data or {}).get("months") or []
    return max((as_int(item.get("month")) for item in months
                if as_int(item.get("count")) > 0), default=0)


def xlsx_candidates(system: str):
    """系統資料夾中所有可用的 .xlsx（排除 Excel 開檔時的 ~$ 暫存檔）。"""
    return sorted((path for path in (DATA_DIR / system).glob("*.xlsx")
                   if not path.name.startswith("~$")), key=lambda path: path.name)


def pick_latest_xlsx(system: str, cache: dict, parser):
    """挑資料夾裡「資料月份最新」的那份統計檔，回傳 (files, best, errors, files_read)。

    原本是挑 mtime 最新的檔，但① 舊檔常留在資料夾裡（620 的「_7月」）、② Google Drive
    同步會重寫 mtime —— 兩者湊在一起，別台機器就可能挑到舊檔，月份卡在 7 月不動。
    改成每個候選檔都解析（有簽章快取，沒更新不會重開 Excel），比誰的資料涵蓋到最新月份。
    """
    best = None
    errors = []
    files_read = 0
    files = xlsx_candidates(system)
    for path in files:
        try:
            data, cache_hit = cached_parse(cache, path, parser)
        except Exception as exc:
            errors.append(f"{path.name}：{exc}")
            continue
        files_read += 0 if cache_hit else 1
        # 月份相同時（例：同月份的兩份匯出）取數字較大、再取檔案較新的那份。
        rank = (latest_data_month(data), as_int((data or {}).get("count")), path.stat().st_mtime)
        if best is None or rank > best[0]:
            best = (rank, path, data, cache_hit)
    return files, best, errors, files_read


def parse_620(path: Path):
    """解析 620 月份追蹤的歷史表，累加 2026-01 起的檢索次數。"""
    if load_workbook is None:
        raise RuntimeError("openpyxl 未安裝，無法讀取 .xlsx")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["月份追蹤"]
        except KeyError:
            raise ValueError("找不到名稱為「月份追蹤」的工作表")
        # 第 1 列的操作說明也提到「歷史表」；只接受真正的區塊標題「▼歷史表」。
        history_row = next((row for row in range(1, sheet.max_row + 1)
                            if str(sheet.cell(row, 1).value or "").strip().startswith("▼歷史表")), None)
        if history_row is None:
            raise ValueError("找不到「歷史表」區塊")
        header_row = history_row + 1
        date_col = next((col for col in range(1, sheet.max_column + 1)
                         if "年/月" in str(sheet.cell(header_row, col).value or "")), None)
        count_col = next((col for col in range(1, sheet.max_column + 1)
                          if "檢索次數" in str(sheet.cell(header_row, col).value or "")), None)
        if date_col is None or count_col is None:
            raise ValueError("歷史表找不到「年/月」或「檢索次數」欄")
        months = []
        # 只巡覽歷史表的「年/月」到「檢索次數」欄位；遇到已開始的歷史資料後首個空月份便停止，
        # 不掃工作表預先保留的數百列空白區。
        first_col, last_col = min(date_col, count_col), max(date_col, count_col)
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, min_col=first_col,
                                                           max_col=last_col, values_only=True), header_row + 1):
            raw_month = as_int(row[date_col - first_col])
            if not raw_month and months:
                break
            if raw_month < 202601 or raw_month > 209912:
                continue
            count = as_int(row[count_col - first_col])
            months.append({"month": raw_month, "count": count, "row": row_number})
        return {"count": sum(item["count"] for item in months), "months": months}
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def system_620(cache: dict):
    """包裝 620 解析結果；來源檔沒更新時直接使用 JSON 快取。"""
    result = {
        "name": "620", "count": 0, "files_found": 0, "files_read": 0,
        "file": "", "months": [], "errors": [], "cache_hit": False, "updated_month": 0,
        "source": "月份追蹤工作表歷史表（2026-01 起）",
    }
    files, best, errors, files_read = pick_latest_xlsx("620", cache, parse_620)
    result.update(files_found=len(files), files_read=files_read, errors=errors)
    if best is None:
        if not errors:
            result["errors"].append("找不到 .xlsx 統計檔")
        return result
    _rank, path, data, cache_hit = best
    result.update(data or {}, file=path.name, cache_hit=cache_hit,
                  updated_month=latest_data_month(data))
    return result


def parse_630_file(path: Path):
    """讀取單一 630 月檔的「全部」工作表最下方「合計」列。"""
    if xlrd is None:
        raise RuntimeError("xlrd 未安裝，無法讀取舊式 .xls")
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        try:
            sheet = workbook.sheet_by_name("全部")
        except xlrd.biffh.XLRDError:
            raise ValueError("找不到名稱為「全部」的工作表")
        count_col = next((col for col in range(sheet.ncols)
                          if "檢索次數" in str(sheet.cell_value(0, col) or "")), None)
        if count_col is None:
            raise ValueError("找不到「檢索次數」欄")
        # 每日明細下方已有 Excel 產生的「合計」列；只取它，避免再把每日值加一次。
        total_row = max((row for row in range(1, sheet.nrows)
                         if str(sheet.cell_value(row, 0) or "").strip() == "合計"), default=None)
        if total_row is None:
            raise ValueError("找不到「合計」列")
        return as_int(sheet.cell_value(total_row, count_col)), total_row + 1
    finally:
        workbook.release_resources()


def system_630(cache: dict):
    """逐月比對檔案簽章，只解析新檔或有更新的 630 月檔。"""
    files = sorted((DATA_DIR / "630").glob("*.xls"), key=lambda p: (month_key(p), p.name))
    total = 0
    months, errors = [], []
    files_read = 0
    cache_hits = 0
    for path in files:
        try:
            data, cache_hit = cached_parse(cache, path, parse_630_file)
            count, row = data
            total += count
            months.append({"file": path.name, "month": month_key(path), "count": count, "row": row})
            files_read += 0 if cache_hit else 1
            cache_hits += 1 if cache_hit else 0
        except Exception as exc:
            errors.append(f"{path.name}：{exc}")
    # 630 月檔檔名只有月份（例：「(8月)營運狀況表.xls」），年度統計固定看當年度。
    latest = max((item["month"] for item in months
                  if 1 <= item["month"] <= 12 and item["count"] > 0), default=0)
    return {
        "name": "630", "count": total, "files_found": len(files), "files_read": files_read,
        "cache_hits": cache_hits, "months": months, "errors": errors,
        "updated_month": date.today().year * 100 + latest if latest else 0,
        "source": "各月營運狀況表「全部」工作表的「合計」列",
    }


def parse_680(path: Path):
    """解析 680 統計總表：取「累計」區塊的「總計」列，並順便記下每個月份的檢索次數。

    表格是「每月一個區塊（月份寫在 A 欄、區塊內 B 欄有『總計』列），最後接一個『累計』區塊」。
    整年 12 個月的區塊是先排好的，還沒發生的月份數字為空或 0 —— 有數字的最後一個月份才是
    這份檔案真正更新到的月份，前端的「更新至」就是看它。
    """
    if load_workbook is None:
        raise RuntimeError("openpyxl 未安裝，無法讀取 .xlsx")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["統計表"]
        except KeyError:
            raise ValueError("找不到名稱為「統計表」的工作表")
        count_col = None
        months = {}
        month_rows = {}
        current_month = 0
        seen_cumulative = False
        total_row = None
        # 單趟循序掃描：表頭前 5 列找「檢索次數」欄，之後邊走邊分辨月份區塊與累計區塊。
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if count_col is None:
                if row_number > 5:
                    break
                count_col = next((index for index, value in enumerate(row, 1)
                                  if "檢索次數" in str(value or "")), None)
                continue
            first = as_int(row[0]) if row else 0
            first_text = str((row[0] if row else None) or "").strip()
            second = str((row[1] if len(row) > 1 else None) or "").strip()
            value = as_int(row[count_col - 1]) if len(row) >= count_col else 0
            if first_text == "累計":
                seen_cumulative = True
            elif 202601 <= first <= 209912:
                current_month = first
            if second != "總計":
                continue
            if seen_cumulative:
                total_row = (row_number, value)      # 累計區塊可能有兩列總計，取最後一列
            elif current_month:
                # 每個月份區塊也有兩列數值相同的總計（依系統版本／依國別），取較大者即可。
                months[current_month] = max(months.get(current_month, 0), value)
                month_rows.setdefault(current_month, row_number)
        if count_col is None:
            raise ValueError("找不到「檢索次數」欄")
        if not seen_cumulative or total_row is None:
            raise ValueError("找不到累計區塊的「總計」列")
        return {
            "count": total_row[1], "row": total_row[0],
            # 還沒發生的月份區塊是空的（0），不列進來，「更新至」才不會被預排的月份拉走。
            "months": [{"month": month, "count": months[month], "row": month_rows[month]}
                       for month in sorted(months) if months[month] > 0],
        }
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def system_680(cache: dict):
    """包裝 680 解析結果；來源總表沒更新時直接使用 JSON 快取。"""
    result = {
        "name": "680", "count": 0, "files_found": 0, "files_read": 0,
        "file": "", "months": [], "errors": [], "cache_hit": False, "updated_month": 0,
        "source": "統計總表最下方的「累計／總計」列",
    }
    files, best, errors, files_read = pick_latest_xlsx("680", cache, parse_680)
    result.update(files_found=len(files), files_read=files_read, errors=errors)
    if best is None:
        if not errors:
            result["errors"].append("找不到 .xlsx 統計總表")
        return result
    _rank, path, data, cache_hit = best
    result.update(data or {}, file=path.name, cache_hit=cache_hit,
                  updated_month=latest_data_month(data))
    return result


def statistics():
    """使用檔案簽章快取統計；僅在使用者按更新時執行，不做背景自動更新。"""
    with _CACHE_LOCK:
        cache = load_cache()
        systems = [
            system_620(cache),
            system_630(cache),
            system_680(cache),
        ]
        try:
            save_cache(cache)
        except OSError:
            pass                     # 快取寫不成不影響本次統計結果
    target = load_annual_target()
    total = sum(int(system.get("count", 0) or 0) for system in systems)
    return {
        "ok": True, "target": target, "total": total,
        "percent": total / target * 100 if target else 0,
        "remaining": max(0, target - total), "systems": systems,
    }


class Handler(BaseHTTPRequestHandler):
    def send_json(self, obj, code=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store, must-revalidate")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = self.path.split("?", 1)[0]
        if path == "/api/stats":
            try:
                self.send_json(statistics())
            except Exception as exc:
                self.send_json({"ok": False, "systems": [], "message": f"統計計算失敗：{exc}"}, 500)
            return
        if path in ("/", "/index.html"):
            target = WEB_DIR / "index.html"
            content_type = "text/html; charset=utf-8"
        elif path == "/style.css":
            target = WEB_DIR / "style.css"
            content_type = "text/css; charset=utf-8"
        elif path == "/app.js":
            target = WEB_DIR / "app.js"
            content_type = "application/javascript; charset=utf-8"
        else:
            self.send_error(404)
            return
        try:
            body = target.read_bytes()
        except OSError:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store, must-revalidate")
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, fmt, *args):
        print(f"[web] {self.address_string()} - {fmt % args}")


def main():
    server = ThreadingHTTPServer(("127.0.0.1", PORT), Handler)
    url = f"http://127.0.0.1:{PORT}"
    print(f"年度檢索統計已啟動：{url}")
    print("按 Ctrl+C 可停止。")
    webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n年度檢索統計已停止。")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
